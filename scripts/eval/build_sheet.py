#!/usr/bin/env python3
"""Add one comparison stage to the golden-set workbook.

Each stage is one change to the tutor (baseline, new prompt, new retrieval, new
STT/TTS), run twice by golden_latency.py -- cold and warm -- and set beside
AFE-Learning-App's benchmark of the same 20 questions.

    python3 scripts/eval/build_sheet.py --stage "1 Baseline" \\
        --cold apps/backend/logs/golden/<run>-cold --warm apps/backend/logs/golden/<run>-warm \\
        --note "Old retrieval, old prompt, the app's own generation settings"

    python3 scripts/eval/build_sheet.py --stage "4 Speech (sherpa)" --speech apps/backend/logs/golden/<run>-speech \\
        --note "..."     # a speech stage, from speech_bench.py

Writes docs/golden-comparison.xlsx: a sheet per stage (re-running a stage replaces
its sheet) and a Summary sheet with one cold row and one warm row per stage plus
AFE's reference row. Needs openpyxl (pip install openpyxl).

AFE's benchmark ran with the model resident and each prompt new, so it compares
with the WARM rows; AFE has no cold numbers.
"""

from __future__ import annotations

import argparse
import json
import statistics
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is needed: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "golden-comparison.xlsx"
AFE_RUNS = ROOT.parent / "AFE-Learning-App" / "dev-data" / "logs" / "benchmark"
PREAMBLE_END = "prioritizing their wording and examples over your own knowledge."

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDD6FE")
COLD_FILL = PatternFill("solid", fgColor="DBEAFE")
WARM_FILL = PatternFill("solid", fgColor="FEF3C7")
AFE_FILL = PatternFill("solid", fgColor="DCFCE7")
WRAP = Alignment(wrap_text=True, vertical="top")

METRICS = [  # (label, poc key, afe getter)
    ("load ms", "load_ms", lambda a: a["llm"]["loadMs"]),
    ("retrieval ms", "retrieval_ms", lambda a: a["retrieval"]["ms"]),
    ("first token ms", "ttft_ms", lambda a: a["llm"]["ttftMs"]),
    ("prefill ms", "prefill_ms", lambda a: a["llm"]["prefillMs"]),
    ("decode ms", "decode_ms", lambda a: a["llm"]["decodeMs"]),
    ("prompt tokens", "prompt_tokens", lambda a: a["llm"]["promptTokens"]),
    ("output tokens", "output_tokens", lambda a: a["llm"]["evalTokens"]),
    ("tokens/s", "tokens_per_second", lambda a: a["llm"]["tokensPerSecond"]),
]


def load_run(path: Path) -> Dict[str, Dict[str, Any]]:
    return {r["id"]: r for r in json.loads((path / "results.json").read_text())}


def load_afe(path: Path):
    results = {r["id"]: r for r in json.loads((path / "results.json").read_text())["results"]}
    prompts = {}
    for line in (path / "prompts.jsonl").read_text().splitlines():
        if line.strip():
            p = json.loads(line)
            prompts[p["turnId"]] = p["messages"]
    return results, prompts


def median(values) -> Optional[float]:
    values = [v for v in values if isinstance(v, (int, float))]
    return round(statistics.median(values), 1) if values else None


def mean(values) -> Optional[float]:
    values = [v for v in values if isinstance(v, (int, float))]
    return round(statistics.mean(values), 2) if values else None


def split_system(system: str):
    head, _, excerpts = system.partition(PREAMBLE_END)
    return head.strip(), " ".join(excerpts.split())


def sources(messages: List[Dict[str, str]]) -> str:
    if not messages:
        return ""
    lines = [ln for ln in messages[0]["content"].splitlines() if ln.startswith("[") and "]" in ln[:5]]
    return "; ".join(ln.split("] ", 1)[-1] for ln in lines)


def summary_rows(stage: str, cold, warm, afe) -> List[List[Any]]:
    ids = [i for i in afe if i in warm]
    rows = []
    for label, run in (("cold", cold), ("warm", warm)):
        if not run:
            continue
        rs = [run[i] for i in ids if i in run]
        rows.append(
            [stage, label, len(rs)]
            + [median(r[k] for r in rs) for _, k, _ in METRICS]
            + [mean(r["answer_facts"] for r in rs), mean(r["context_facts"] for r in rs)]
        )
    a = [afe[i] for i in ids]
    rows.append(
        ["AFE (reference)", "warm", len(a)]
        + [median(g(x) for x in a) for _, _, g in METRICS]
        + [mean(x["answerFactCoverage"] for x in a), mean(x["retrieval"]["factCoverage"] for x in a)]
    )
    return rows


SUMMARY_HEAD = ["stage", "model", "n"] + [m[0] + " (median)" for m in METRICS] + ["answer facts (mean)", "context facts (mean)"]


def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def fill_row(ws, row: int, ncols: int, fill) -> None:
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = fill


def write_stage(wb, stage: str, note: str, cold, warm, cold_prompts, warm_prompts, afe, afe_prompts, afe_dir: Path, warm_dir: Path) -> None:
    title = stage[:31]
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    any_run = warm or cold
    first = next(iter(any_run.values()))
    ws["A1"] = "Stage {}".format(stage)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = note
    ws["A3"] = "Model {}  |  options {}  |  this app: {}  |  AFE: {}".format(
        first.get("model"), json.dumps(first.get("options")), warm_dir.name.rsplit("-", 1)[0], afe_dir.name)
    ws["A4"] = ("cold = LLM and embedding model unloaded before every question (pays the load). "
                "warm = models resident, each prompt new to the cache. AFE's run is warm.")

    r = 6
    for c, h in enumerate(SUMMARY_HEAD, 1):
        ws.cell(row=r, column=c, value=h)
    style_header(ws, r, len(SUMMARY_HEAD))
    for row in summary_rows(stage, cold, warm, afe):
        r += 1
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        fill_row(ws, r, len(SUMMARY_HEAD), AFE_FILL if str(row[0]).startswith("AFE") else
                 (COLD_FILL if row[1] == "cold" else WARM_FILL))

    r += 3
    head = (["id", "type", "question"]
            + ["cold " + m[0] for m in METRICS]
            + ["warm " + m[0] for m in METRICS]
            + ["AFE " + m[0] for m in METRICS]
            + ["answer facts cold", "answer facts warm", "answer facts AFE",
               "context facts (this app)", "context facts AFE",
               "persona same as AFE", "passages same as AFE", "user message same as AFE",
               "sources (this app)", "sources (AFE)",
               "answer (warm)", "answer (AFE)", "expected", "manual: grounded yes/no", "manual: notes"])
    head_row = r
    for c, h in enumerate(head, 1):
        ws.cell(row=r, column=c, value=h)
    style_header(ws, r, len(head))
    for qid, a in afe.items():
        w, c_ = warm.get(qid), cold.get(qid)
        if not (w or c_):
            continue
        ref = w or c_
        pm = (warm_prompts.get(qid) or cold_prompts.get(qid) or {}).get("messages", [])
        am = afe_prompts.get(a["turnId"], [])
        p_persona, p_exc = split_system(pm[0]["content"]) if pm else ("", "")
        a_persona, a_exc = split_system(am[0]["content"]) if am else ("", "")
        r += 1
        values = ([qid, ref["type"], ref["question"]]
                  + [c_.get(k) if c_ else None for _, k, _ in METRICS]
                  + [w.get(k) if w else None for _, k, _ in METRICS]
                  + [g(a) for _, _, g in METRICS]
                  + [c_["answer_facts"] if c_ else None, w["answer_facts"] if w else None, a["answerFactCoverage"],
                     ref["context_facts"], a["retrieval"]["factCoverage"],
                     "yes" if p_persona == a_persona else "no",
                     "yes" if p_exc == a_exc else "no",
                     "yes" if pm and am and pm[-1]["content"] == am[-1]["content"] else "no",
                     sources(pm), sources(am),
                     (w or c_)["answer"], a.get("answer", ""), ref.get("expected"), "", ""])
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.alignment = WRAP
    ws.freeze_panes = ws.cell(row=head_row + 1, column=4)
    widths = {1: 6, 2: 10, 3: 40}
    for col in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 12)
    for col in range(len(head) - 6, len(head) - 1):  # sources, answers, expected
        ws.column_dimensions[get_column_letter(col)].width = 50


def write_summary(wb, stage: str, cold, warm, afe) -> None:
    ws = wb["Summary"] if "Summary" in wb.sheetnames else wb.create_sheet("Summary", 0)
    if ws.max_row < 2 or ws["A1"].value != "Golden set comparison":
        ws.delete_rows(1, ws.max_row)
        ws["A1"] = "Golden set comparison"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = ("One cold and one warm row per stage; AFE's reference row is warm. "
                    "Medians over the 20 questions; facts are means of keyword coverage.")
        for c, h in enumerate(SUMMARY_HEAD, 1):
            ws.cell(row=4, column=c, value=h)
        style_header(ws, 4, len(SUMMARY_HEAD))
    # drop this stage's old rows and the AFE row, then append
    for row in range(ws.max_row, 4, -1):
        if ws.cell(row=row, column=1).value in (stage, "AFE (reference)"):
            ws.delete_rows(row)
    rows = summary_rows(stage, cold, warm, afe)
    for row in rows:  # the AFE row comes last, so it always ends the table
        r = ws.max_row + 1
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        fill_row(ws, r, len(SUMMARY_HEAD), AFE_FILL if str(row[0]).startswith("AFE") else
                 (COLD_FILL if row[1] == "cold" else WARM_FILL))
    ws.column_dimensions["A"].width = 22
    for col in range(2, len(SUMMARY_HEAD) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14


SPEECH_HEAD = ["", "n", "STT WER (mean)", "words", "STT after stop ms (mean)", "STT after stop ms (max)",
               "TTS first sentence ms (mean)", "TTS RTF (mean)", "TTS RTF (max)", "STT model", "TTS voice"]


def write_speech(wb, stage: str, note: str, speech_dir: Path, afe, afe_doc, afe_dir: Path) -> None:
    """A speech stage: speech_bench.py's STT/TTS numbers beside AFE's, per question."""
    run = json.loads((speech_dir / "results.json").read_text())
    ours = {r["id"]: r for r in run["results"]}
    summ = run["summary"]
    title = stage[:31]
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws["A1"] = "Stage {}".format(stage)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = note
    ws["A3"] = "This app: {}  |  AFE: {}".format(speech_dir.name, afe_dir.name)
    ws["A4"] = ("Audio is each question spoken by the app's own TTS voice (AFE does the same), so WER is "
                "optimistic. 'After stop' = wait between the mic closing and the final text. Models resident; "
                "cold loads are in the backend's startup log.")
    r = 6
    for c, h in enumerate(SPEECH_HEAD, 1):
        ws.cell(row=r, column=c, value=h)
    style_header(ws, r, len(SPEECH_HEAD))
    a_stt = [x["stt"] for x in afe.values() if x.get("stt", {}).get("wer") is not None]
    a_tts = [x["tts"] for x in afe.values() if x.get("tts", {}).get("rtf") is not None]
    rows = [
        ["this app", len(ours), summ["wer_mean"], summ["words"], summ["after_stop_ms_mean"], summ["after_stop_ms_max"],
         summ["first_sentence_ms_mean"], summ["rtf_mean"], summ["rtf_max"], summ["stt_model"], summ["tts_voice"]],
        ["AFE (reference)", len(a_stt), round(statistics.mean(s["wer"] for s in a_stt), 3), sum(s["words"] for s in a_stt),
         round(statistics.mean(s["afterStopMs"] for s in a_stt), 1), max(s["afterStopMs"] for s in a_stt),
         round(statistics.mean(t["firstSentenceMs"] for t in a_tts), 1),
         round(statistics.mean(t["rtf"] for t in a_tts), 3), max(t["rtf"] for t in a_tts),
         a_stt[0]["model"] if a_stt else "", a_tts[0]["voice"] if a_tts else ""],
    ]
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        fill_row(ws, r, len(SPEECH_HEAD), AFE_FILL if row[0].startswith("AFE") else WARM_FILL)

    r += 3
    head = ["id", "type", "question", "audio s", "WER", "WER AFE", "after stop ms", "after stop ms AFE",
            "transcript", "transcript AFE", "TTS sentences", "TTS sentences AFE", "first sentence ms",
            "first sentence ms AFE", "synth ms", "synth ms AFE", "audio s", "audio s AFE", "RTF", "RTF AFE"]
    head_row = r
    for c, h in enumerate(head, 1):
        ws.cell(row=r, column=c, value=h)
    style_header(ws, r, len(head))
    for qid, a in afe.items():
        o = ours.get(qid)
        if not o:
            continue
        ast, att, ost, ott = a.get("stt", {}), a.get("tts", {}), o["stt"], o["tts"]
        r += 1
        values = [qid, o["type"], o["question"], ost["audio_sec"], ost["wer"], ast.get("wer"),
                  ost["after_stop_ms"], ast.get("afterStopMs"), ost["transcript"], ast.get("transcript"),
                  ott.get("sentences"), att.get("sentences"), ott.get("first_sentence_ms"), att.get("firstSentenceMs"),
                  ott.get("synth_ms"), att.get("synthMs"), ott.get("audio_sec"), att.get("audioSec"),
                  ott.get("rtf"), att.get("rtf")]
        for c, v in enumerate(values, 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
    ws.freeze_panes = ws.cell(row=head_row + 1, column=4)
    for col in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(col)].width = {1: 6, 2: 10, 3: 40, 9: 40, 10: 40}.get(col, 13)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--stage", required=True)
    ap.add_argument("--note", default="")
    ap.add_argument("--cold")
    ap.add_argument("--warm")
    ap.add_argument("--speech", help="a speech_bench.py run dir: writes a speech stage instead")
    ap.add_argument("--afe", help="AFE benchmark run dir (default: newest)")
    ap.add_argument("--out", default=str(OUT))
    args = ap.parse_args()

    afe_dir = Path(args.afe) if args.afe else sorted(p for p in AFE_RUNS.iterdir() if p.is_dir())[-1]
    afe, afe_prompts = load_afe(afe_dir)
    out = Path(args.out)
    if args.speech:
        wb = load_workbook(out) if out.exists() else Workbook()
        write_speech(wb, args.stage, args.note, Path(args.speech), afe, None, afe_dir)
        wb.save(out)
        print("Wrote speech stage '{}' to {}".format(args.stage, out))
        return 0
    if not args.warm:
        sys.exit("--warm (or --speech) is required")
    warm_dir = Path(args.warm)
    warm = load_run(warm_dir)
    warm_prompts = {p["id"]: p for p in json.loads((warm_dir / "prompts.json").read_text())}
    cold, cold_prompts = {}, {}
    if args.cold:
        cold = load_run(Path(args.cold))
        cold_prompts = {p["id"]: p for p in json.loads((Path(args.cold) / "prompts.json").read_text())}

    wb = load_workbook(out) if out.exists() else Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        wb["Sheet"].title = "Summary"
    write_summary(wb, args.stage, cold, warm, afe)
    write_stage(wb, args.stage, args.note, cold, warm, cold_prompts, warm_prompts, afe, afe_prompts, afe_dir, warm_dir)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("Wrote stage '{}' to {}".format(args.stage, out))
    for row in summary_rows(args.stage, cold, warm, afe):
        print("  " + " | ".join(str(v) for v in row))
    return 0


if __name__ == "__main__":
    sys.exit(main())
